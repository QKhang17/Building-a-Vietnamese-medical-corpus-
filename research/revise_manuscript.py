"""Tao ban thao phuong phap benchmark, khong ghi de ban goc va khong bia so lieu."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "NCKH_ban_rut_gon_15_trang.docx"
OUTPUT = ROOT / "output" / "NCKH_ban_benchmark_300_draft.docx"
FIGURE_1 = ROOT / "research" / "figures" / "hinh-1-pipeline-lai.png"
FIGURE_2 = ROOT / "research" / "figures" / "hinh-2-gan-nhan-phan-xu.png"


def set_text(paragraph, text: str) -> None:
    paragraph.clear()
    paragraph.add_run(text)


def remove_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)


def shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def add_table_after(
    document: Document,
    paragraph,
    caption: str,
    headers: list[str],
    rows: list[list[str]],
):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    for index, value in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = value
        shade(cell, "D9EAF7")
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(7.5)
                    run.bold = row_index == 0
    caption_paragraph = document.add_paragraph(caption)
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in caption_paragraph.runs:
        run.italic = True
        run.font.name = "Times New Roman"
        run.font.size = Pt(9)
    paragraph._p.addnext(caption_paragraph._p)
    caption_paragraph._p.addnext(table._tbl)
    return table


def insert_figure_before(paragraph, image_path: Path, caption: str) -> None:
    figure_paragraph = paragraph.insert_paragraph_before()
    figure_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure_paragraph.paragraph_format.page_break_before = False
    figure_paragraph.paragraph_format.keep_with_next = True
    figure_paragraph.paragraph_format.space_before = Pt(0)
    figure_paragraph.paragraph_format.space_after = Pt(0)
    figure_paragraph.add_run().add_picture(str(image_path), width=Inches(6.75))
    caption_paragraph = paragraph.insert_paragraph_before(caption)
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.page_break_before = False
    caption_paragraph.paragraph_format.keep_with_next = True
    caption_paragraph.paragraph_format.space_before = Pt(0)
    caption_paragraph.paragraph_format.space_after = Pt(3)
    for run in caption_paragraph.runs:
        run.italic = True
        run.font.name = "Times New Roman"
        run.font.size = Pt(9)


def main() -> None:
    document = Document(SOURCE)
    original = list(document.paragraphs)

    # Canh bao ro rang de ban draft khong bi nham la ket qua da hoan tat.
    status = original[1].insert_paragraph_before(
        "GHI CHÚ TRẠNG THÁI: Bản này đã cập nhật phương pháp và hạ tầng benchmark; "
        "chưa có 300 nhãn vàng hoặc kết quả ablation nên chưa dùng để nộp."
    )
    status.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in status.runs:
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x9B, 0x1C, 0x1C)

    set_text(
        original[1],
        "Tóm tắt—Nghiên cứu trình bày MedNLP Studio, một nền tảng xây dựng và kiểm chứng "
        "dữ liệu thực thể y khoa tiếng Việt theo quy trình human-in-the-loop có truy vết. "
        "Hệ thống kết hợp từ điển có phiên bản với Gemini, kiểm tra schema, loại, offset và "
        "điều kiện surface = X[start:end], đồng thời ghi log từng ứng viên, token và độ trễ. "
        "Một giao thức khóa trước chọn phân tầng 300 tóm tắt từ kho 4.182 bài, chia 30 pilot, "
        "70 development và 200 test; hai chuyên gia gán độc lập và người thứ ba phân xử. "
        "Bốn cấu hình ablation, bootstrap theo tài liệu và nghiên cứu đối trọng blind–assisted "
        "được định nghĩa trước khi mở test. Tại ngày khóa bản phương pháp, mẫu 300 tài liệu và "
        "900 assignment đã được tạo nhưng nghiên cứu con người và benchmark chưa hoàn tất; vì "
        "vậy bài không báo cáo trước P/R/F1 hoặc lợi ích thời gian."
    )
    set_text(
        original[7],
        "Nghiên cứu đóng góp ba sản phẩm có thể kiểm chứng sau khi hoàn tất: (i) corpus 300 "
        "tóm tắt có hai bản gán độc lập và một snapshot vàng; (ii) benchmark ablation bốn cấu "
        "hình trên test khóa; và (iii) nền tảng annotation có version, audit event, optimistic "
        "locking và trace từng ứng viên AI. Bản hiện tại cung cấp đầy đủ hạ tầng và giao thức, "
        "không coi kế hoạch hoặc số assignment là bằng chứng về độ chính xác NER."
    )

    set_text(
        original[11],
        "Rà soát chốt ngày 30/08/2026 xác nhận sáu corpus NER y–sinh có thành phần tiếng Việt "
        "công khai theo tiêu chí đã định: PhoNER_COVID19, VietBioNER, ViMQ, ViMedNER, "
        "VietMed-NER và En‑ViMedNER [18]–[23]. Chúng khác nhau về nguồn (tin dịch tễ, tài "
        "liệu lao, câu hỏi bệnh nhân, website sức khỏe, lời nói và bản dịch PubMed), nên không "
        "thể thay thế trực tiếp một corpus tóm tắt khoa học tiếng Việt bản địa có mã thuật ngữ. "
        "Về công cụ, đối thủ gần nhất là Prodigy và INCEpTION chứ không chỉ các mô hình NER: "
        "cả hai đã hỗ trợ pre-annotation và review/curation nhiều người [24], [25]. Vì vậy tính "
        "mới của MedNLP Studio được giới hạn ở gói tích hợp chuyên biệt và benchmark truy vết, "
        "không tuyên bố phát minh human-in-the-loop."
    )
    add_table_after(
        document,
        original[11],
        "Bảng 2. So sánh MedNLP Studio với các nền tảng và engine liên quan.",
        ["Nền tảng", "Việt y khoa", "Mã", "Pre/LLM", "Trace/audit", "Độc lập + phân xử"],
        [
            ["MedNLP Studio", "Có", "Có", "Dictionary + Gemini ràng buộc", "Run/event/version/hash", "Có, 2+1"],
            ["Prodigy", "Tùy biến", "Recipe", "Model/pattern/LLM", "Session/dataset", "Có, review"],
            ["INCEpTION", "Tùy biến", "Knowledge base", "Recommender", "Project/document", "Có, curation"],
            ["doccano", "Tùy biến", "Không nêu", "API/import", "Project/user", "Collaborative; adjudication không nêu"],
            ["brat", "Tùy biến", "Attribute", "Tùy biến", "Có thể log từng edit", "Cộng tác thời gian thực"],
            ["cTAKES", "Không", "UMLS", "Engine NLP", "Do người dùng xây", "Không áp dụng"],
            ["MedTagger", "Không", "OMOP", "Dictionary/rule/ML", "Do người dùng xây", "Không áp dụng"],
        ],
    )
    add_table_after(
        document,
        original[11],
        "Bảng 1. Các corpus NER y–sinh tiếng Việt được xác nhận tại ngày chốt.",
        ["Corpus", "Nguồn", "Quy mô công bố", "Đặc trưng", "Khoảng trống so với nghiên cứu này"],
        [
            ["PhoNER_COVID19", "Tin COVID", "10K câu/35K entity", "10 loại", "Dịch tễ, không mã"],
            ["VietBioNER", "Tài liệu lao", "1.706 câu", "5 loại, brat", "Miền hẹp"],
            ["ViMQ", "Câu hỏi bệnh nhân", "9.000 câu", "3 loại + intent", "Hội thoại"],
            ["ViMedNER", "4 website", "Theo repository", "5 nhóm chẩn đoán/điều trị", "Không phải tóm tắt học thuật"],
            ["VietMed-NER", "Lời nói", "18 loại", "Spoken NER", "Có nhiễu ASR"],
            ["En‑ViMedNER", "PubMed dịch", "4.392 cặp abstract", "21 UMLS types", "Dẫn xuất cross-lingual"],
        ],
    )

    set_text(
        original[18],
        "Pipeline nhận tóm tắt nguyên văn cùng article ID và checksum. Nhánh từ điển sinh "
        "span/type/code và giữ manifest hash; nhánh Gemini sinh JSON term/start/end/type/code "
        "với model, temperature và prompt hash đã khóa. Từng ứng viên được phân loại thành "
        "accepted, invalid JSON/schema/type/offset, surface mismatch, duplicate, overlap removed "
        "hoặc dictionary override. Trace gồm token, latency từng giai đoạn và raw response cache."
    )
    original[19].clear()
    original[19].alignment = WD_ALIGN_PARAGRAPH.CENTER
    original[19].paragraph_format.page_break_before = False
    original[19].paragraph_format.keep_with_next = True
    original[19].paragraph_format.space_before = Pt(0)
    original[19].paragraph_format.space_after = Pt(0)
    original[19].add_run().add_picture(str(FIGURE_1), width=Inches(6.75))
    set_text(original[20], "Hình 1. Pipeline lai có ràng buộc và truy vết thực nghiệm.")
    original[20].alignment = WD_ALIGN_PARAGRAPH.CENTER
    original[20].paragraph_format.page_break_before = False
    original[20].paragraph_format.keep_with_next = True
    original[20].paragraph_format.space_before = Pt(0)
    original[20].paragraph_format.space_after = Pt(3)
    set_text(
        original[21],
        "Hình 1 tách rõ dữ liệu đầu vào, hai nhánh sinh ứng viên, lớp hậu kiểm, hợp nhất và "
        "kho trace. Từ điển được chạy cả trong cấu hình độc lập và hybrid; AI raw/constrained "
        "không nhận ngữ cảnh từ điển. Chỉ hybrid cho phép metadata từ điển ghi đè ứng viên AI "
        "trùng span. Nhờ lưu raw response, mọi phép chấm lại dùng cùng đầu ra mô hình."
    )

    # Bo hinh cu ve tu dien; so do quy trinh 2+1 duoc dat o thiet ke kiem chung.
    remove_paragraph(original[28])
    remove_paragraph(original[29])
    set_text(
        original[30],
        "Kiểm tra mã nằm trước bước lọc mục mơ hồ, bởi một bí danh chỉ có giá trị khi liên kết "
        "được với khái niệm hợp lệ. Manifest ghi phiên bản, số mục và mã băm; kết quả cũ vì vậy "
        "vẫn xác định được tài nguyên đã dùng. Độ phủ không được suy ra từ kích thước từ điển mà "
        "được đo bằng exact-span recall trên gold, tách theo sáu loại cùng unique-surface "
        "coverage và valid-code rate."
    )

    set_text(
        original[36],
        "Dữ liệu nghiên cứu được tách thành project, document, assignment, entity, audit event "
        "và gold snapshot. Entity có start, end, surface, type, code, source, decision và version. "
        "Ghi entity dùng optimistic locking trên entities_version; mọi lần lưu/undo/nộp/khóa "
        "đều có before–after audit. API nhận xét cũ được giữ để tương thích nhưng không được "
        "dùng thay cho annotation có cấu trúc."
    )
    set_text(
        original[37],
        "Giao diện mới có blind annotation và assisted review: chọn/sửa span, đổi loại/mã, "
        "thêm/xóa, accept/reject, lý do và hoàn tác. Luồng trạng thái là assigned → in_progress "
        "→ submitted → conflict → adjudicated → locked. Blind không nhận pre-label; annotator "
        "không thấy bản của nhau. Adjudicator chỉ mở sau conflict, xem hai bản song song và tạo "
        "một snapshot vàng bất biến."
    )

    set_text(
        original[42],
        "Từ 4.182 bản ghi, hệ thống loại abstract rỗng/lỗi và khử trùng bằng SHA-256 trên văn "
        "bản NFKC, chuẩn hóa khoảng trắng và casefold. Mẫu 300 được phân tầng theo domain nguồn, "
        "năm và tam phân vị độ dài với seed 20260830; manifest công bố đúng 30 pilot, 70 "
        "development và 200 test. Test bị khóa trước khi sửa prompt, luật hoặc từ điển."
    )
    set_text(
        original[43],
        "Hai chuyên gia gán độc lập sáu loại; người thứ ba phân xử. Trên development, thiết kế "
        "đối trọng phân 35 mẫu A-assisted/B-blind và 35 mẫu A-blind/B-assisted. Hệ thống đo "
        "active time, accepted/modified/rejected/added và exact quality so với gold. Agreement "
        "trước phân xử được báo cáo bằng exact-span P/R/F1 theo loại và micro/macro."
    )
    set_text(
        original[44],
        "Test chạy bốn cấu hình: dictionary; AI raw; AI constrained; hybrid. Ba cấu hình AI "
        "chạy ba lần; 10 development abstracts được dùng warm-up trước khi đo 200 test. Báo cáo "
        "exact micro/macro và theo loại, relaxed overlap, code accuracy, ma trận 6×6, taxonomy "
        "lỗi, rejection rate, dictionary coverage và latency/cost. Khoảng tin cậy và chênh lệch "
        "ghép cặp dùng bootstrap theo tài liệu 1.000 lần."
    )
    insert_figure_before(
        original[45],
        FIGURE_2,
        "Hình 2. Gán độc lập, nghiên cứu hỗ trợ và phân xử.",
    )

    set_text(
        original[46],
        "Hạ tầng đã vượt 75 kiểm thử tự động bao phủ offset Unicode, validation, state "
        "transition, chống lộ pre-label ở blind, điều kiện mở phân xử, sampling tái lập, metric "
        "tính tay, bootstrap, rejection logging và hợp nhất lai. Front-end build thành công và "
        "workspace localhost tải không có lỗi console. Đây là kết quả phần mềm, không phải kết "
        "quả độ chính xác NER."
    )
    set_text(
        original[47],
        "Project mednlp-gold-v1 đã được tạo trong MySQL với 300 document và 900 assignment: "
        "600 assignment annotator (530 blind, 70 assisted) và 300 adjudicator. Ở ngày khóa bản "
        "này có 0/300 gold snapshot; 69/70 assisted assignment còn chờ preannotation Gemini. "
        "Các assignment đó bị API từ chối mở cho đến khi snapshot tiền nhãn được tạo."
    )
    set_text(
        original[48],
        "Chưa chạy benchmark test hoặc nghiên cứu con người; vì thế các bảng P/R/F1, agreement, "
        "rejection rate, confusion, coverage, latency, cost và thời gian chuyên gia chưa có giá "
        "trị. Script xuất bản chính thức sẽ dừng nếu thiếu 300 gold snapshot hoặc thiếu 10 run "
        "files (1 dictionary + 3×3 AI), nhằm ngăn dữ liệu chưa hoàn chỉnh lọt vào bài."
    )
    add_table_after(
        document,
        original[48],
        "Bảng 3. Trạng thái triển khai tại ngày khóa bản phương pháp.",
        ["Hạng mục", "Yêu cầu", "Hiện có", "Trạng thái"],
        [
            ["Mẫu phân tầng", "300 = 30/70/200", "300 = 30/70/200", "Đạt"],
            ["Assignment", "2 annotator + 1 adjudicator/tài liệu", "900", "Đạt"],
            ["Preannotation development", "70", "1", "Chờ 69 lượt Gemini"],
            ["Gold snapshot", "300", "0", "Chờ chuyên gia"],
            ["Ablation test", "1 D + 3 R + 3 C + 3 H", "0", "Chờ gold và ngân sách API"],
            ["Bảng kết quả", "Sinh từ raw trace", "Schema/script sẵn sàng", "Chưa sinh số liệu"],
        ],
    )

    set_text(
        original[50],
        "Ưu điểm đã kiểm chứng của phiên bản mới là khả năng truy vết và cách ly vai trò, không "
        "phải F1 chưa đo. Mỗi span nối được về văn bản/checksum, assignment/version, quyết định "
        "chuyên gia, dictionary manifest hoặc candidate event của AI. Raw cache giúp tính lại "
        "metric mà không phát sinh lời gọi mô hình mới; script báo cáo giữ mẫu số rõ ràng cho "
        "từng nguyên nhân loại ứng viên."
    )
    set_text(
        original[51],
        "Giới hạn chính hiện nay là nghiên cứu con người chưa diễn ra và chưa có số liệu test. "
        "Không thể suy ra hiệu quả của hybrid, lợi ích thời gian hay độ phủ từ điển từ test phần "
        "mềm. Schema v1.0 cũng không hỗ trợ span gián đoạn/lồng nhau; sáu loại có thể cần hiệu "
        "chỉnh sau pilot; chi phí và biến thiên dịch vụ Gemini phải được chốt tại ngày chạy."
    )
    set_text(
        original[54],
        "MedNLP Studio đã được nâng từ prototype nhận xét theo bài thành nền tảng annotation có "
        "entity-level editing, gán độc lập, phân xử, audit và benchmark trace. Mẫu 300 tóm tắt "
        "đã khóa bằng seed/checksum, nhưng đóng góp corpus và kết luận về hiệu quả chỉ tồn tại "
        "sau khi đủ 300 snapshot vàng và các run test hoàn tất."
    )
    set_text(
        original[55],
        "Bước tiếp theo bắt buộc là chạy 69 preannotation development theo bảng giá chính thức, "
        "thực hiện pilot, khóa guideline, hoàn tất gán/phân xử, rồi mới chạy ablation. Sau đó "
        "các CSV được sinh tự động sẽ thay thế bảng trạng thái trong bản draft. Nếu AI được dùng "
        "để hỗ trợ ngôn ngữ, nhóm tác giả phải công bố theo chính sách nơi nộp và chịu trách "
        "nhiệm toàn bộ nội dung."
    )

    set_text(
        original[73],
        '[17] Google, “Gemini 2.5 Flash,” Gemini API Documentation, https://ai.google.dev/gemini-api/docs/models, accessed Aug. 30, 2026.'
    )

    for reference in [
        '[18] VinAI Research, “PhoNER_COVID19,” https://github.com/VinAIResearch/PhoNER_COVID19, accessed Aug. 30, 2026.',
        '[19] U. Phan et al., “A Named Entity Recognition Corpus for Vietnamese Biomedical Texts to Support Tuberculosis Treatment,” LREC, 2022.',
        '[20] T. D. Huy et al., “ViMQ: A Vietnamese Medical Question Dataset for Healthcare Dialogue System Development,” ICONIP, 2021/2023.',
        '[21] P. V. Duong et al., “ViMedNER: A Medical Named Entity Recognition Dataset for Vietnamese,” EAI INIS, 2024.',
        '[22] K. Le-Duc et al., “Medical Spoken Named Entity Recognition,” NAACL Industry Track, 2025.',
        '[23] N. Vo et al., “En-ViMedNER: An English–Vietnamese Parallel Biomedical Corpus with UMLS Semantic Type Annotations,” EMNLP, to appear, 2026.',
        '[24] Explosion, “Prodigy Built-in Recipes and Review,” https://prodi.gy/docs/recipes, accessed Aug. 30, 2026.',
        '[25] INCEpTION Project, “INCEpTION User Guide: Annotation and Curation,” https://inception-project.github.io/documentation/latest/user-guide, accessed Aug. 30, 2026.',
        '[26] Google, “Gemini Developer API Pricing,” https://ai.google.dev/gemini-api/docs/pricing, accessed Aug. 30, 2026.',
        '[27] OHNLP, “MedTagger,” https://github.com/OHNLP/MedTagger, accessed Aug. 30, 2026.',
    ]:
        paragraph = document.add_paragraph(reference)
        paragraph.style = original[57].style
        paragraph.paragraph_format.page_break_before = False
        paragraph.paragraph_format.keep_with_next = False
        paragraph.paragraph_format.keep_together = False
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        for run in paragraph.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(9)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
